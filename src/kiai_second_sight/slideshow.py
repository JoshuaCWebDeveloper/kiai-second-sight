from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt


def rebuild_slideshow(manifest: dict, output: str | Path) -> Path:
    output = Path(output)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(10)
    blank = prs.slide_layouts[6]

    for card in manifest.get("cards", []):
        question = card.get("question_image")
        answer = card.get("answer_image")
        if not question or not answer:
            continue
        for kind, image_path in (("Question", question), ("Answer", answer)):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
            if kind == "Answer":
                box = slide.shapes.add_textbox(Inches(0.22), Inches(9.55), Inches(9.55), Inches(0.32))
                p = box.text_frame.paragraphs[0]
                p.text = (
                    f"Move {card['move_number']} · played {card['played_move']} · "
                    f"{card['winrate_before']*100:.1f}% → {card['winrate_after']*100:.1f}% "
                    f"(-{card['loss_pp']*100:.1f} pp)"
                )
                p.font.size = Pt(12)

    # python-pptx creates a default title slide in some templates only when explicitly added;
    # with our blank-only path there are no extra slides.
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output
